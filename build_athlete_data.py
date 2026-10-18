# -*- coding: utf-8 -*-
"""يستخرج بيانات نظام الرياضيين (مراحل، أنواع أيام، وجبات، أغذية لكل 100 جم، إرشادات) من Diet_Pro_Athlete.xlsx إلى src/athlete.json"""
import json, os, sys
from openpyxl import load_workbook
H = os.path.dirname(os.path.abspath(__file__)); SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(H, '..', 'Diet_Pro_Athlete.xlsx')
wb = load_workbook(SRC, data_only=True)
num = lambda v, d=0.0: float(v) if isinstance(v, (int, float)) else d
st = wb['Pro إعدادات']
phases = [{'n': st.cell(r, 1).value, 'p': num(st.cell(r, 2).value), 'f': num(st.cell(r, 3).value), 'adj': num(st.cell(r, 4).value), 'note': st.cell(r, 5).value or ''} for r in range(5, 11) if st.cell(r, 1).value]
dayTypes = [{'n': st.cell(r, 1).value, 'carb': num(st.cell(r, 2).value), 'kpm': num(st.cell(r, 3).value), 'train': st.cell(r, 4).value == 'نعم'} for r in range(14, 20) if st.cell(r, 1).value]
fac = {'neat': num(st.cell(22, 2).value, 1.3), 'minP': num(st.cell(23, 2).value, 1.6), 'minF': num(st.cell(24, 2).value, .6), 'minEA': num(st.cell(25, 2).value, 30), 'water': num(st.cell(26, 2).value, 35), 'waterHr': num(st.cell(27, 2).value, 500), 'tol': num(st.cell(28, 2).value, 5), 'minC': num(st.cell(29, 2).value, 2)}
meals = [{'n': st.cell(r, 1).value, 't': str(st.cell(r, 2).value or ''), 'on': st.cell(r, 3).value == 'نعم', 'rest': [num(st.cell(r, c).value) for c in (4, 5, 6)], 'train': [num(st.cell(r, c).value) for c in (7, 8, 9)]} for r in range(33, 39) if st.cell(r, 1).value]
db = wb['Pro قاعدة الأغذية']; foods = []
for r in range(5, db.max_row + 1):
    if not db.cell(r, 3).value or db.cell(r, 6).value is None: continue
    foods.append({'id': db.cell(r, 1).value, 'g': db.cell(r, 2).value, 'n': db.cell(r, 3).value, 'st': db.cell(r, 5).value or '', 'k': num(db.cell(r, 6).value), 'p': num(db.cell(r, 7).value), 'c': num(db.cell(r, 8).value), 'f': num(db.cell(r, 9).value), 'fi': num(db.cell(r, 10).value), 'su': num(db.cell(r, 11).value), 'na': num(db.cell(r, 12).value)})
# قياسات منزلية للأغذية الأساسية (ورقة «أغذية أساسية للرياضيين» في Pro_Foods_Database.xlsx) — تُطابَق بالاسم العربي
PF = os.path.join(H, '..', 'Pro_Foods_Database.xlsx'); ms = {}
if os.path.exists(PF):
    w2 = load_workbook(PF, data_only=True, read_only=True)['أغذية أساسية للرياضيين']
    for r in w2.iter_rows(min_row=2, values_only=True):
        if r[1] and r[11]: ms[str(r[1]).strip()] = str(r[11])
for f in foods:
    if f['n'] in ms: f['ms'] = ms[f['n']]
gd = wb['Pro إرشادات المراحل']; guides = [[gd.cell(r, 1).value, gd.cell(r, 2).value, gd.cell(r, 3).value] for r in range(5, gd.max_row + 1) if gd.cell(r, 3).value and gd.cell(r, 4).value == 'نعم']
pf = wb['Pro تفضيلات']; prefs = {'maxProt': num(pf.cell(4, 2).value, 5), 'maxAny': num(pf.cell(5, 2).value, 8), 'maxItems': num(pf.cell(6, 2).value, 4), 'veg': num(pf.cell(7, 2).value, 150), 'egg': num(pf.cell(8, 2).value, 50)}
out = {'phases': phases, 'dayTypes': dayTypes, 'fac': fac, 'meals': meals, 'foods': foods, 'guides': guides, 'prefs': prefs, 'days': ['الأحد', 'الاثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة', 'السبت']}
json.dump(out, open(os.path.join(H, 'src', 'athlete.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
from collections import Counter
print('قياسات منزلية:', sum(1 for f in foods if 'ms' in f), 'من', len(foods)); print('athlete.json:', len(phases), 'مراحل |', len(dayTypes), 'أنواع أيام |', len(meals), 'وجبات |', len(foods), 'صنفاً |', len(guides), 'إرشاداً |', round(os.path.getsize(os.path.join(H, 'src', 'athlete.json')) / 1024), 'KB')
print('المجموعات:', dict(Counter(f['g'] for f in foods)))
print('أنواع الإرشادات:', dict(Counter(g[1] for g in guides)), '| مراحل:', sorted(set(g[0] for g in guides)))
