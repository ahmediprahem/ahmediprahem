# -*- coding: utf-8 -*-
"""يستخرج بيانات النظام من Diet_Pro_Max.xlsx إلى src/data.json (قراءة فقط): الأصناف (مع معايرة الكميات)، بنك الوجبات، القواعد السريرية،
بعده شغّل: node tag_foods.js (وسوم الأصناف + الأصناف والوجبات المصرية/السعودية)
الحساسيات، جداول الإعدادات، وقوالب التركيب الأسبوعية. التشغيل:  py build_data.py [المسار\\Diet_Pro_Max.xlsx]"""
import json, os, sys
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
EX = [(15, 3, 0, 80), (15, 7, 3, 115), (0, 7, 1, 35), (0, 7, 3, 55), (0, 7, 5, 75), (0, 7, 8, 100), (15, 8, 3, 120), (15, 8, 5, 135), (15, 8, 8, 165), (5, 2, 0, 25), (15, 0, 0, 60), (0, 0, 5, 45), (15, 0, 0, 60)]
GROUPS = ['نشويات', 'بقوليات', 'لحوم شديدة الانخفاض بالدهون', 'لحوم قليلة الدهون', 'لحوم متوسطة الدهون', 'لحوم عالية الدهون', 'حليب خالي الدسم', 'حليب قليل الدسم', 'حليب كامل الدسم', 'خضروات غير نشوية', 'فواكه', 'دهون', 'كربوهيدرات أخرى']
KEY = [1, 1, 2, 2, 2, 2, 2, 2, 2, 1, 1, 3, 1]          # الماكرو الأساسي لمعايرة كمية كل مجموعة (1=كارب 2=بروتين 3=دهون)
# ---- قوالب التركيب الأسبوعية (أوزان، سقوف، حد أدنى) لكل مجموعة × فتحة (فطور، س1، غداء، س2، عشاء، س3) — نفس افتراضيات ورقة «تركيبة الأسبوع»
BASE = {0: [(3, 2), (0, 0), (4, 4), (0, 0), (3, 2), (0, 0)], 1: [(1, 1), (0, 0), (2, 2), (0, 0), (1, 1), (0, 0)], 2: [(1, 1), (0, 0), (4, 4), (0, 0), (2, 2), (0, 0)], 3: [(1, 1), (0, 0), (4, 4), (0, 0), (2, 2), (0, 0)],
        4: [(3, 2), (0, 0), (1, 1), (0, 0), (2, 2), (0, 0)], 5: [(0, 0), (0, 0), (1, 2), (0, 0), (0, 1), (0, 0)], 6: [(1, 1), (1, 1), (0, 0), (0, 0), (1, 1), (0, 0)], 7: [(1, 1), (1, 1), (0, 0), (0, 0), (1, 1), (0, 0)],
        8: [(1, 1), (1, 1), (0, 0), (0, 0), (1, 1), (0, 0)], 9: [(1, 1), (0, 0), (3, 3), (0, 0), (3, 3), (0, 0)], 10: [(1, 1), (2, 1), (1, 1), (3, 2), (0, 0), (0, 0)], 11: [(1, 1), (2, 2), (3, 3), (2, 2), (2, 2), (0, 0)], 12: [(0, 0)] * 6}
DINNER = {1: {0: (3, 1), 2: (1, 1), 3: (2, 2), 4: (2, 2), 6: (0, 0), 7: (0, 0), 8: (0, 0), 10: (0, 0), 11: (2, 1)}, 2: {0: (3, 2), 2: (2, 2), 3: (1, 1), 4: (1, 1), 6: (0, 0), 7: (0, 0), 8: (0, 0), 10: (0, 0), 11: (2, 1)},
          3: {0: (0, 0), 1: (0, 0), 2: (0, 0), 3: (2, 1), 4: (1, 1), 6: (3, 1), 7: (3, 1), 8: (3, 1), 9: (2, 2), 10: (3, 1), 11: (3, 1)}, 4: {0: (3, 1), 2: (0, 0), 3: (3, 2), 4: (2, 1), 6: (1, 1), 7: (1, 1), 8: (1, 1), 10: (0, 0), 11: (2, 1)},
          5: {0: (0, 0), 1: (0, 0), 2: (0, 0), 3: (2, 1), 4: (3, 2), 6: (2, 1), 7: (2, 1), 8: (2, 1), 10: (2, 1), 11: (2, 1)}, 6: {0: (3, 2), 1: (2, 1), 2: (1, 1), 3: (1, 1), 4: (1, 1), 6: (0, 0), 7: (0, 0), 8: (0, 0), 10: (0, 0), 11: (2, 1)},
          7: {0: (4, 4), 2: (3, 3), 3: (2, 2), 9: (3, 3), 11: (3, 2)}}
LUNCH7 = {0: (2, 2), 2: (2, 2), 3: (2, 2), 11: (2, 2)}; SWEET_SLOT = {1: 1, 2: 3, 3: 1, 4: 3, 5: 3, 6: 3, 7: 3}
def comp_default(d):
    out = []
    for g in range(13):
        row = [[BASE[g][s][0], BASE[g][s][1], 0] for s in range(6)]
        if g in DINNER[d]: row[4][0], row[4][1] = DINNER[d][g]
        if d == 7 and g in LUNCH7: row[2][0], row[2][1] = LUNCH7[g]
        if g == 12: row[SWEET_SLOT[d]][0], row[SWEET_SLOT[d]][1] = 3, 1
        if g == 11: row[SWEET_SLOT[d]][2] = 1
        out.append(row)
    return out
def num(v, d=0):
    return v if isinstance(v, (int, float)) else d
def build(xlsx):
    wb = load_workbook(xlsx, data_only=True); db, mb, st, cn, pf, wd, sp = wb['قاعدة البدائل'], wb['بنك الوجبات'], wb['الإعدادات'], wb['القواعد السريرية'], wb['تفضيلات العميل'], wb['بيانات Word'], wb['عبارات الدعم']
    foods = []
    for r in range(5, 1005):
        nm = db.cell(r, 3).value
        if not nm: continue
        g = GROUPS.index(db.cell(r, 2).value) if db.cell(r, 2).value in GROUPS else None
        if g is None: continue
        exv = num(db.cell(r, 9).value, 1) or 1; mac = [num(db.cell(r, c).value) for c in (10, 11, 12, 13)]
        act = mac[KEY[g] - 1] / exv; cal = 1.0
        if act > 0: cal = max(0.6, min(2.0, EX[g][KEY[g] - 1] / act))
        units = num(db.cell(r, 5).value, 1); grams = num(db.cell(r, 8).value)
        foods.append({'n': str(nm), 'g': g, 'us': db.cell(r, 6).value or '', 'up': db.cell(r, 7).value or '', 'upe': round(units / exv * cal, 4), 'gpe': round(grams / exv * cal, 3), 'units': units, 'exv': exv, 'mac': mac,
                      'cooked': db.cell(r, 18).value or '', 'm': [1 if str(db.cell(r, c).value) == 'نعم' else 0 for c in (19, 20, 21, 22)], 'veg': 1 if str(db.cell(r, 23).value) == 'نعم' else 0, 'vegan': 1 if str(db.cell(r, 24).value) == 'نعم' else 0,
                      'on': 0 if str(db.cell(r, 26).value) == 'لا' else 1, 'gi': db.cell(r, 48).value if isinstance(db.cell(r, 48).value, (int, float)) else None, 'o': r})
    bank = []
    for r in range(6, 166):
        nm = mb.cell(r, 2).value
        if not nm: continue
        ing = []
        for c in (13, 16, 19, 22, 25, 28):
            if mb.cell(r, c).value: ing.append({'n': str(mb.cell(r, c).value), 'e': num(mb.cell(r, c + 1).value, 1), 'role': str(mb.cell(r, c + 2).value or '')})
        bank.append({'id': str(mb.cell(r, 1).value), 'n': str(nm), 'm': [1 if str(mb.cell(r, c).value) == 'نعم' else 0 for c in (3, 4, 5, 6)], 'scal': 1 if str(mb.cell(r, 7).value) == 'نعم' else 0, 'max': num(mb.cell(r, 8).value, 4), 'min': num(mb.cell(r, 9).value),
                     'on': 0 if str(mb.cell(r, 12).value) == 'لا' else 1, 'ing': ing, 'note': str(mb.cell(r, 31).value or ''), 'time': str(mb.cell(r, 32).value or '')})
    rules = [{'kw': str(cn.cell(r, 1).value), 'cond': str(cn.cell(r, 2).value), 'tier': str(cn.cell(r, 3).value), 'why': str(cn.cell(r, 4).value or ''), 'src': str(cn.cell(r, 5).value or '')} for r in range(13, 313) if cn.cell(r, 1).value and str(cn.cell(r, 7).value) == 'نعم']
    instr = [{'cond': str(cn.cell(r, 2).value), 'cat': str(cn.cell(r, 3).value or ''), 'text': str(cn.cell(r, 4).value), 'src': str(cn.cell(r, 5).value or '')} for r in range(317, 417) if cn.cell(r, 4).value and str(cn.cell(r, 6).value) == 'نعم']
    allergens = [{'label': str(pf.cell(r, 20).value), 'kw': [k.strip() for k in str(pf.cell(r, 21).value).replace('،', ',').split(',') if k.strip()], 'type': str(pf.cell(r, 22).value or 'حساسية')} for r in range(6, 31) if pf.cell(r, 20).value and pf.cell(r, 21).value]
    pct = {}
    for r in range(52, 56): pct[str(int(st.cell(r, 1).value))] = [num(st.cell(r, c).value) for c in range(2, 8)]
    data = {'version': 1, 'EX': EX, 'groups': GROUPS, 'foods': foods, 'bank': bank, 'rules': rules, 'instr': instr, 'allergens': allergens,
            'settings': {'activity': [[str(st.cell(r, 1).value), num(st.cell(r, 2).value)] for r in range(5, 10)], 'goals': [[str(st.cell(r, 1).value), num(st.cell(r, 2).value)] for r in range(12, 18)],
                         'ibw_off': num(st.cell(5, 4).value, 100), 'ibw_f': num(st.cell(6, 4).value, 0.9), 'adj': num(st.cell(7, 4).value, 0.4), 'water': num(st.cell(8, 4).value, 35), 'min_kcal': num(st.cell(12, 4).value, 1200),
                         'min_pro': num(st.cell(13, 4).value, 1.2), 'max_pro': num(st.cell(14, 4).value, 2.5), 'max_def': num(st.cell(15, 4).value, 0.25), 'tol': [10, 10, 5], 'pct': pct, 'types': ['B', 'S', 'L', 'S', 'D', 'S'], 'slotNames': ['الإفطار', 'سناك 1', 'الغداء', 'سناك 2', 'العشاء', 'سناك 3']},
            'comp': [comp_default(d) for d in range(1, 8)], 'days': ['الأحد', 'الاثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة', 'السبت'],
            'tips': [[str(wd.cell(r, 2).value or ''), str(wd.cell(r, 3).value or '')] for r in range(98, 106)], 'phrases': [[str(sp.cell(r, 1).value), str(sp.cell(r, 2).value)] for r in range(5, 261) if sp.cell(r, 2).value and str(sp.cell(r, 3).value) != 'لا'], 'times': ['07:30', '10:30', '13:30', '16:30', '20:00', '22:00'], 'conditions': ['خمول الغدة', 'سكري النوع الثاني', 'مقاومة الإنسولين / ما قبل السكري']}
    return data
if __name__ == '__main__':
    x = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'Diet_Pro_Max.xlsx'); d = build(x)
    json.dump(d, open(os.path.join(HERE, 'src', 'data.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print('data.json:', len(d['foods']), 'صنفاً |', len(d['bank']), 'وجبة بنك |', len(d['rules']), 'قاعدة سريرية |', len(d['instr']), 'تعليمة |', len(d['allergens']), 'حساسية |', round(os.path.getsize(os.path.join(HERE, 'src', 'data.json')) / 1024), 'KB')
